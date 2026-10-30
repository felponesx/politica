"""Segunda rodada das fontes novas: formato real de TCU (dicionário/API), IBAMA, lista suja, Receita (xlsx), INPE anual e CAGED (FTP). Saída curta."""
import ftplib
import io
import re

import httpx

from scripts.diagnostico_ci import zip_remoto

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36", "Accept": "*/*"}
cli = httpx.Client(timeout=120, follow_redirects=True, headers=UA)


def topo(nome, url, n=900, rng=None):
    try:
        h = dict(UA)
        if rng:
            h["Range"] = f"bytes=0-{rng}"
        r = cli.get(url, headers=h)
        print(f"\n== {nome}: {r.status_code} {r.headers.get('content-type','')[:30]} {r.headers.get('content-length')}B", flush=True)
        if r.status_code in (200, 206):
            print(r.content[:n].decode("latin-1", "replace").replace("\r", ""), flush=True)
        return r
    except Exception as e:
        print(f"\n== {nome}: ERRO {type(e).__name__} {str(e)[:100]}", flush=True)


r = cli.get("https://sites.tcu.gov.br/dados-abertos/inidoneos-irregulares/dicionario-dados.html")
txt = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", r.text))
i = txt.lower().find("irregular")
print("TCU dicionário:", r.status_code, txt[max(0, i - 100):i + 900], flush=True)
for u in ("https://certidoes.apps.tcu.gov.br/api/publico/responsaveis-contas-irregulares?offset=0&limit=2",
          "https://certidoes.apps.tcu.gov.br/api/publico/responsaveis-contas-irregulares/"):
    topo("TCU " + u[-45:], u, 300)
topo("IBAMA bioma.csv", "https://stibamadadosabertosprd.blob.core.windows.net/dados-abertos/dados/SIFISC/auto_infracao/bioma/bioma.csv", 500, rng=3000)
zip_remoto("IBAMA auto_infracao_csv.zip", "https://stibamadadosabertosprd.blob.core.windows.net/dados-abertos/dados/SIFISC/auto_infracao/auto_infracao/auto_infracao_csv.zip", mostrar=("",), linhas=3)
topo("Lista suja CSV", "https://www.gov.br/trabalho-e-emprego/pt-br/assuntos/inspecao-do-trabalho/areas-de-atuacao/cadastro_de_empregadores.csv", 700, rng=4000)
try:
    import openpyxl
    r = cli.get("https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/publicacoes/relatorios/renuncia/gastos-tributarios-bases-efetivas/dgt-bases-efetivas-2023-serie-2021-a-2026-quadros.xlsx")
    print("\n== Receita xlsx:", r.status_code, len(r.content), flush=True)
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    print("abas:", wb.sheetnames[:20], flush=True)
    for ws in wb.worksheets[:3]:
        print("--", ws.title, flush=True)
        for k, lin in enumerate(ws.iter_rows(values_only=True)):
            if k > 7:
                break
            print([c for c in lin if c is not None][:7], flush=True)
except Exception as e:
    print("Receita xlsx ERRO", type(e).__name__, str(e)[:120], flush=True)
r = cli.get("https://dataserver-coids.inpe.br/queimadas/queimadas/focos/csv/anual/")
print("\n== INPE anual:", r.status_code, re.findall(r'href="([^"?/][^"]*)"', r.text)[-8:], flush=True)
try:
    f = ftplib.FTP(timeout=30)
    f.connect("ftp.mtps.gov.br", 21)
    f.login()
    f.cwd("/pdet/microdados/NOVO CAGED")
    anos = f.nlst()
    print("\n== CAGED NOVO CAGED:", anos[-6:], flush=True)
    f.cwd(anos[-1])
    meses = f.nlst()
    print("   ", anos[-1], meses[-4:], flush=True)
    f.cwd(meses[-1])
    for nome in f.nlst()[:6]:
        try:
            print("    ", nome, f.size(nome), flush=True)
        except Exception:
            print("    ", nome, flush=True)
    f.quit()
except Exception as e:
    print("CAGED ERRO", type(e).__name__, str(e)[:120], flush=True)
print("fim", flush=True)

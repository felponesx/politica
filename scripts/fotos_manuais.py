"""Fotos de pessoas dos escândalos conferidas à mão (planilha do Felipe), guardadas no repositório.

  python -m scripts.fotos_manuais importar <planilha.xlsx>   # uma vez: tira as fotos da aba "Galeria de fotos" para app/dados/fotos_manuais/
  python -m scripts.fotos_manuais aplicar [site/midia]       # coloca nos blocos de fotos (chave `esc:<nome no caso>`) e grava creditos_fotos.json

`fotos_manuais()` devolve {"esc:<nome>": data-URL} e é usada por `scripts/midia_site.py`, para que refazer as fotos não apague estas. A foto manual vale sobre a
automática. Foto de grupo ou montagem (a pessoa é uma entre várias) não entra: no avatar redondo apareceria o rosto errado (`PULAR`).
"""
import base64
import io
import json
import logging
import os
import re
import sys
import unicodedata

PASTA = os.path.join(os.path.dirname(__file__), "..", "app", "dados", "fotos_manuais")
INDICE = os.path.join(PASTA, "fotos_manuais.json")
log = logging.getLogger("brasilplus.midia")
# linhas da galeria com foto de grupo, montagem ou imagem com texto por cima: ficam fora até haver um recorte só do rosto
PULAR = {6, 28, 35, 36, 37, 38, 46, 47}


def _slug(t: str) -> str:
    t = unicodedata.normalize("NFD", t).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")[:60]


def importar(xlsx: str):
    import openpyxl
    from PIL import Image
    from app import escandalos
    wb = openpyxl.load_workbook(xlsx)
    g, s = wb["Galeria de fotos"], wb["Pessoas sem foto"]
    nomes = {p["nome"] for c in escandalos.casos().values() for p in c["pessoas"]}
    os.makedirs(PASTA, exist_ok=True)
    idx, pulados = {}, []
    for im in g._images:
        r = im.anchor._from.row + 1
        linha = g.cell(r, 8).value
        nome = s.cell(int(linha), 2).value if linha else g.cell(r, 2).value  # nome como está no caso (a aba "Pessoas sem foto")
        if r in PULAR:
            pulados.append(nome)
            continue
        if nome not in nomes:
            log.warning("sem correspondência no caso: %s", nome)
            continue
        arq = _slug(nome) + ".webp"
        Image.open(io.BytesIO(im._data())).convert("RGB").save(os.path.join(PASTA, arq), "WEBP", quality=90)
        idx[nome] = {"arquivo": arq, "fonte": g.cell(r, 6).value, "licenca": g.cell(r, 7).value, "observacao": g.cell(r, 5).value}
    with open(INDICE, "w", encoding="utf-8") as f:
        json.dump(idx, f, ensure_ascii=False, indent=1)
    print(f"{len(idx)} fotos importadas; fora (grupo/montagem): {pulados}")


def fotos_manuais() -> dict:
    try:
        idx = json.load(open(INDICE, encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    out = {}
    for nome, d in idx.items():
        try:
            with open(os.path.join(PASTA, d["arquivo"]), "rb") as f:
                out["esc:" + nome] = "data:image/webp;base64," + base64.b64encode(f.read()).decode()
        except OSError:
            log.warning("foto manual sem arquivo: %s", nome)
    return out


def creditos() -> dict:
    try:
        idx = json.load(open(INDICE, encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return {n: {"fonte": d.get("fonte"), "licenca": d.get("licenca")} for n, d in idx.items()}


def aplicar(pasta: str = "site/midia"):
    from scripts.midia_site import fnv1a
    idx = json.load(open(os.path.join(pasta, "indice.json"), encoding="utf-8"))
    n = idx["blocos"]
    fotos = fotos_manuais()
    blocos = {i: json.load(open(os.path.join(pasta, f"fotos_{i}.json"), encoding="utf-8")) for i in range(n)}
    for k, v in fotos.items():
        blocos[fnv1a(k) % n][k] = v
    for i, b in blocos.items():
        with open(os.path.join(pasta, f"fotos_{i}.json"), "w", encoding="utf-8") as f:
            json.dump(b, f, separators=(",", ":"))
    with open(os.path.join(pasta, "creditos_fotos.json"), "w", encoding="utf-8") as f:
        json.dump(creditos(), f, ensure_ascii=False, separators=(",", ":"))
    idx["fotos"] = sum(len(b) for b in blocos.values())
    with open(os.path.join(pasta, "indice.json"), "w", encoding="utf-8") as f:
        json.dump(idx, f, ensure_ascii=False)
    print(f"{len(fotos)} fotos manuais aplicadas")


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    if len(sys.argv) > 2 and sys.argv[1] == "importar":
        importar(sys.argv[2])
    elif sys.argv[1:2] == ["aplicar"]:
        aplicar(*sys.argv[2:3])
    else:
        print(__doc__)

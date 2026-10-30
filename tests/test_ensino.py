"""Ensino: leitura das planilhas do IDEB (INEP) e a rota /ensino."""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.collectors import inep
from app.db import SessionLocal, criar_tabelas
from app.main import app
from app.models import Registro


def _planilha(titulo, cab, linhas, obs=(2021, 2023), proj=(2021,)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo
    for _ in range(9):
        ws.append([None] * 5)
    ws.append(cab + [f"VL_OBSERVADO_{a}" for a in obs] + [f"VL_PROJECAO_{a}" for a in proj])
    for l in linhas:
        ws.append(l)
    return wb


def _salva(wb, tmp_path, nome="x.xlsx"):
    caminho = tmp_path / nome
    wb.save(caminho)
    return caminho


def test_linhas_da_aba_acha_a_linha_de_codigos_e_trata_traco(tmp_path):
    wb = _planilha("UF e Regiões (AI)", ["Região/UF", "Rede"], [["Acre", "Total (3)", 5.1, "-", 6.0], ["Norte", "Pública", 4.0, 4.4, 5.0]])
    wb = inep.abrir_xlsx(_salva(wb, tmp_path))
    out = list(inep.linhas_da_aba(wb.worksheets[0], 2))
    assert out[0] == (("Acre", "Total (3)"), {2021: 5.1, 2023: None}, {2021: 6.0})
    assert inep._sem_nulos(out[0][1]) == {"2021": 5.1}
    assert inep._tira_nota("Privada (1)") == "Privada"


def test_ler_ufs_pega_so_a_rede_total_e_ignora_regioes(monkeypatch, tmp_path):
    def aba(sigla):
        return _planilha(f"UF e Regiões ({sigla})", ["Região/UF", "Rede"], [["Norte", "Total", 4.0, 4.5, 5.0], ["Acre", "Total (3)", 5.0, 5.2, 5.5], ["Acre", "Privada", 7.0, 7.1, 7.0],
                                                                         ["Distrito Federal", "Total", 6.0, 6.2, 6.5], ["Brasil", "Total", 5.5, 5.8, 6.0]])
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sigla in ("AI", "AF", "EM"):
        ws = wb.create_sheet(f"UF e Regiões ({sigla})")
        for l in aba(sigla).active.iter_rows(values_only=True):
            ws.append(list(l))
    caminho = _salva(wb, tmp_path, "ufs.xlsx")
    monkeypatch.setattr(inep, "baixar", lambda nome: caminho)
    ufs, brasil, metas = inep.ler_ufs(2023)
    assert ufs["iniciais"] == {"AC": {"2021": 5.0, "2023": 5.2}, "DF": {"2021": 6.0, "2023": 6.2}}
    assert brasil["medio"] == {"2021": 5.5, "2023": 5.8} and metas["finais"] == {"2021": 6.0}
    d = inep.montar_ideb(ufs, brasil, metas, 2023)
    assert d["ufs"]["AC"]["iniciais"]["2023"] == 5.2 and d["brasil"]["iniciais"]["2021"] == 5.5 and d["metas"]["medio"] == {"2021": 6.0}


def test_ler_escolas_fica_com_as_melhores_de_cada_estado(monkeypatch, tmp_path):
    linhas = [["SP", 3550308, "São Paulo", 1000 + i, f"ESCOLA {i}", "Estadual", 4.0 + i / 100, 5.0 + i / 100, 6.0] for i in range(100)]
    linhas += [["AC", 1200401, "Rio Branco", 2000, "ESCOLA DO ACRE", "Municipal", 3.0, 3.5, 6.0], ["AC", 1200401, "Rio Branco", 2001, "SEM NOTA", "Municipal", 3.0, "-", 6.0]]
    wb = _planilha("IDEB_Escolas", ["UF", "Cod mun", "Município", "Cod escola", "Escola", "Rede"], linhas)
    caminho = _salva(wb, tmp_path, "esc.xlsx")
    monkeypatch.setattr(inep, "baixar", lambda nome: caminho)
    monkeypatch.setitem(inep.MINIMO_ESCOLAS, "iniciais", 10)
    r = inep.ler_escolas("iniciais", 2023)
    sp = [e for e in r if e["uf"] == "SP"]
    assert len(sp) == inep.ESCOLAS_POR_UF and sp[0]["ideb"] == 6.0 and sp[0]["nome"] == "Escola 95"  # 5,95 a 5,99 arredondam para 6,0: desempata pelo nome
    assert [e["nome"] for e in r if e["uf"] == "AC"] == ["Escola do Acre"]  # sem nota fica de fora
    assert r == sorted(r, key=lambda x: (-x["ideb"], x["nome"]))
    monkeypatch.setitem(inep.MINIMO_ESCOLAS, "iniciais", 1000)
    with pytest.raises(RuntimeError, match="nada foi gravado"):
        inep.ler_escolas("iniciais", 2023)


def test_rota_ensino():
    criar_tabelas()
    cliente = TestClient(app)
    with SessionLocal() as db:
        db.query(Registro).filter(Registro.fonte == "inep").delete()
        db.commit()
    from app.utils import _MEMO
    _MEMO.clear()
    assert cliente.get("/ensino").status_code == 404
    with SessionLocal() as db:
        db.add(Registro(fonte="inep", tipo="ensino_ideb", chave="BR", dados={"ano": 2025, "brasil": {"iniciais": {"2023": 5.2}}, "ufs": {"SP": {"iniciais": {"2023": 6.1}}}, "metas": {}}))
        db.commit()
    try:
        _MEMO.clear()
        d = cliente.get("/ensino").json()
        assert d["ideb"]["ufs"]["SP"]["iniciais"]["2023"] == 6.1 and d["escolas"] is None
    finally:
        with SessionLocal() as db:
            db.query(Registro).filter(Registro.fonte == "inep").delete()
            db.commit()
        _MEMO.clear()


def test_nome_escola_mantem_siglas_e_minusculas_pequenas():
    assert inep.nome_escola("EEEFM EURIDICE LOPES PEDROSO") == "EEEFM Euridice Lopes Pedroso"
    assert inep.nome_escola("COLEGIO ESTADUAL DOUTOR JOSE BALDUINO DE SOUZA DECIO") == "Colegio Estadual Doutor Jose Balduino de Souza Decio"
    assert inep.nome_escola("CED CASA GRANDE") == "CED Casa Grande" and inep.nome_escola("Escola Já Escrita") == "Escola Já Escrita"


def test_ufs_abreviadas_na_planilha(monkeypatch, tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    from app.utils import NOME_PARA_UF
    for sigla in ("AI", "AF", "EM"):
        ws = wb.create_sheet(f"UF e Regiões ({sigla})")
        for _ in range(9):
            ws.append([None] * 4)
        ws.append(["Região/UF", "Rede", "VL_OBSERVADO_2025", "VL_PROJECAO_2021"])
        for nome in list(NOME_PARA_UF) + ["R. G. DO NORTE", "R. G. DO SUL", "M. G. DO SUL"]:
            ws.append([nome.title(), "Total", 5.0, 6.0])
    caminho = _salva(wb, tmp_path, "ab.xlsx")
    monkeypatch.setattr(inep, "baixar", lambda nome: caminho)
    ufs, _, _ = inep.ler_ufs(2025)
    assert {"RN", "RS", "MS"} <= set(ufs["iniciais"])
